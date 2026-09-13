import io

import pytest
from docx import Document as DocxDocument
from fastapi import HTTPException, UploadFile

from app.services import document_processor


def _make_upload_file(data: bytes, content_type: str, filename: str = "resume.pdf") -> UploadFile:
    return UploadFile(filename=filename, file=io.BytesIO(data), headers={"content-type": content_type})


@pytest.mark.asyncio
async def test_validate_rejects_empty_file():
    upload = _make_upload_file(b"", "application/pdf")
    with pytest.raises(HTTPException) as exc:
        await document_processor.validate_and_read_upload(upload)
    assert exc.value.status_code == 400


@pytest.mark.asyncio
async def test_validate_rejects_oversized_file(monkeypatch):
    monkeypatch.setattr(document_processor.settings, "MAX_UPLOAD_MB", 1)
    oversized = b"%PDF-" + b"0" * (2 * 1024 * 1024)
    upload = _make_upload_file(oversized, "application/pdf")
    with pytest.raises(HTTPException) as exc:
        await document_processor.validate_and_read_upload(upload)
    assert exc.value.status_code == 400
    assert "exceeds" in exc.value.detail


@pytest.mark.asyncio
async def test_validate_rejects_disallowed_content_type():
    upload = _make_upload_file(b"%PDF-fake", "text/plain")
    with pytest.raises(HTTPException) as exc:
        await document_processor.validate_and_read_upload(upload)
    assert exc.value.status_code == 400


@pytest.mark.asyncio
async def test_validate_rejects_mismatched_magic_bytes():
    # Content-Type claims PDF but the bytes don't match the PDF or DOCX magic signature.
    upload = _make_upload_file(b"this is not actually a pdf", "application/pdf")
    with pytest.raises(HTTPException) as exc:
        await document_processor.validate_and_read_upload(upload)
    assert exc.value.status_code == 400
    assert "does not match" in exc.value.detail


@pytest.mark.asyncio
async def test_validate_accepts_genuine_docx():
    buf = io.BytesIO()
    doc = DocxDocument()
    doc.add_paragraph("Experienced backend engineer with 5 years of Python.")
    doc.save(buf)
    data = buf.getvalue()

    upload = _make_upload_file(
        data,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        filename="resume.docx",
    )
    result_bytes, kind = await document_processor.validate_and_read_upload(upload)
    assert kind == "docx"
    text = document_processor.extract_text(result_bytes, kind)
    assert "backend engineer" in text.lower()


def test_extract_text_raises_on_empty_docx():
    buf = io.BytesIO()
    DocxDocument().save(buf)  # no paragraphs added
    with pytest.raises(document_processor.DocumentProcessingError):
        document_processor.extract_text(buf.getvalue(), "docx")


def test_extract_text_unsupported_kind_raises():
    with pytest.raises(document_processor.DocumentProcessingError):
        document_processor.extract_text(b"whatever", "txt")
