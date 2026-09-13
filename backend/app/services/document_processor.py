import io

from fastapi import HTTPException, UploadFile

from app.core.config import settings

PDF_MAGIC = b"%PDF-"
DOCX_MAGIC = b"PK\x03\x04"  # DOCX is a zip archive

ALLOWED_CONTENT_TYPES = {
    "application/pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


class DocumentProcessingError(Exception):
    pass


def _sniff_type(data: bytes) -> str:
    if data.startswith(PDF_MAGIC):
        return "pdf"
    if data.startswith(DOCX_MAGIC):
        return "docx"
    raise DocumentProcessingError("File content does not match a supported PDF or DOCX file.")


async def validate_and_read_upload(file: UploadFile) -> tuple[bytes, str]:
    """Validate size + content type (by magic bytes, not just extension) and return (bytes, kind)."""
    max_bytes = settings.MAX_UPLOAD_MB * 1024 * 1024

    data = await file.read()
    if len(data) == 0:
        raise HTTPException(status_code=400, detail="Uploaded file is empty.")
    if len(data) > max_bytes:
        raise HTTPException(status_code=400, detail=f"File exceeds the {settings.MAX_UPLOAD_MB}MB limit.")

    if file.content_type not in ALLOWED_CONTENT_TYPES:
        raise HTTPException(status_code=400, detail="Only PDF and DOCX files are supported.")

    try:
        kind = _sniff_type(data)
    except DocumentProcessingError as e:
        raise HTTPException(status_code=400, detail=str(e))

    return data, kind


def extract_text(data: bytes, kind: str) -> str:
    if kind == "pdf":
        return _extract_pdf(data)
    if kind == "docx":
        return _extract_docx(data)
    raise DocumentProcessingError(f"Unsupported document kind: {kind}")


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        pages = [page.extract_text() or "" for page in reader.pages]
    except Exception as e:
        raise DocumentProcessingError(f"Failed to parse PDF: {e}")

    text = "\n".join(pages).strip()
    if not text:
        raise DocumentProcessingError(
            "No extractable text found in this PDF (it may be a scanned image without OCR)."
        )
    return text


def _extract_docx(data: bytes) -> str:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
        paragraphs = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                paragraphs.extend(cell.text for cell in row.cells)
    except Exception as e:
        raise DocumentProcessingError(f"Failed to parse DOCX: {e}")

    text = "\n".join(p for p in paragraphs if p.strip()).strip()
    if not text:
        raise DocumentProcessingError("No extractable text found in this DOCX file.")
    return text
